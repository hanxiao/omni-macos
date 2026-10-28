# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx", "openpyxl", "numbers-parser"]
# ///
"""The eval bed's decks, workbooks and Numbers documents (bed.json "offices").

    uv run Tools/eval/office.py

Writes Tools/eval/office/, which is committed: `omni-verify evalbed` copies these files into the
tree, so the bed stays the same bytes from run to run. Run again only to change them (the tree hash
changes with them, and every cached eval index is rebuilt).

No speaker notes on the decks: Quick Look draws no preview of a python-pptx deck that has them.
Picture-only slides use built-in user pictures the bed has nowhere else, so a query about one can
only be answered by the slide's picture.
"""
import csv
import datetime
import random
import subprocess
import tempfile
from pathlib import Path

from numbers_parser import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).parent / "office"
STAMP = datetime.datetime(2026, 10, 1, 9, 0, 0)


def png(heic: str, tmp: Path) -> Path:
    out = tmp / (Path(heic).stem + ".png")
    subprocess.run(["sips", "-s", "format", "png", heic, "--out", str(out)], check=True, capture_output=True)
    return out


def deck(name: str, slides: list, tmp: Path) -> None:
    p = Presentation()
    for s in slides:
        if "picture" in s:
            slide = p.slides.add_slide(p.slide_layouts[6])
            slide.shapes.add_picture(str(png(s["picture"], tmp)), Inches(2.5), Inches(1.25), height=Inches(5))
            continue
        slide = p.slides.add_slide(p.slide_layouts[1])
        slide.shapes.title.text = s["title"]
        body = slide.placeholders[1].text_frame
        for i, line in enumerate(s["bullets"]):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = line
    p.core_properties.created = STAMP
    p.core_properties.modified = STAMP
    p.save(OUT / name)


def workbook(name: str, sheets: dict, chart: str | None = None) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, datetime.date):
                    c.number_format = "yyyy-mm-dd"
        if title == chart:
            bars = BarChart()
            bars.title = title
            bars.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(rows)), titles_from_data=True)
            bars.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(rows)))
            ws.add_chart(bars, "D2")
    wb.properties.created = STAMP
    wb.properties.modified = STAMP
    wb.save(OUT / name)


def numbers(name: str, sheets: dict) -> None:
    doc = Document()
    first = True
    for sheet, (table, rows) in sheets.items():
        if first:
            doc.sheets[0].name = sheet
            doc.sheets[0].tables[0].name = table
            first = False
        else:
            doc.add_sheet(sheet, table)
        t = doc.sheets[sheet].tables[0]
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.write(r, c, v)
    doc.save(OUT / name)


def inventory() -> list:
    rng = random.Random(7)
    parts = ["hex bolt", "washer", "bearing", "gasket", "hinge", "spring", "valve", "pulley", "bracket", "clamp"]
    sizes = ["M4", "M6", "M8", "M10", "12 mm", "20 mm", "small", "large"]
    rows = [["SKU", "Part", "Bin", "Qty"]]
    # Items named once among 700 rows of look-alike parts: a row the sheet's other pieces do not echo.
    needles = {95: ["SKU-71190", "Brass ship's bell", "D-31", 4], 260: ["SKU-52406", "Carbon fibre tripod", "A-08", 12],
               511: ["SKU-48213", "Titanium camping kettle", "B-17", 37], 640: ["SKU-66075", "Cast iron dutch oven", "E-22", 9]}
    for i in range(1, 701):
        if i in needles:
            rows.append(needles[i])
            continue
        rows.append([f"SKU-{30000 + rng.randrange(60000)}", f"{rng.choice(parts)} {rng.choice(sizes)}",
                     f"{rng.choice('ACDEF')}-{rng.randrange(1, 40):02d}", rng.randrange(0, 500)])
    return rows


def orders() -> None:
    """A CSV export of 900 look-alike orders with four named once (a record a line: record i is line i + 1)."""
    rng = random.Random(11)
    names = ["Anna", "Ben", "Clara", "David", "Eva", "Felix", "Greta", "Hannes", "Julia", "Karl"]
    family = ["Schmidt", "Meyer", "Wagner", "Becker", "Hoffmann", "Koch", "Richter", "Klein"]
    cities = ["Berlin", "Hamburg", "Munich", "Cologne", "Leipzig", "Dresden", "Bremen", "Hanover"]
    goods = ["USB cable", "notebook", "desk lamp", "phone case", "mouse pad", "water bottle", "backpack", "headphones"]
    needles = {140: ("Ingrid Solberg", "Tromso", "Hand-painted porcelain teapot"),
               455: ("Mateus Ferreira", "Porto", "Vintage brass telescope"),
               701: ("Aiko Tanaka", "Sapporo", "Cedar wood sauna bucket")}
    with open(OUT / "orders.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Order", "Date", "Customer", "City", "Product", "Qty", "Amount EUR"])
        for i in range(1, 901):
            order = "ORD-77310" if i == 860 else f"ORD-{40000 + rng.randrange(30000)}"
            day = datetime.date(2026, 1, 1) + datetime.timedelta(days=rng.randrange(300))
            who, city, what = needles.get(i, (f"{rng.choice(names)} {rng.choice(family)}", rng.choice(cities), rng.choice(goods)))
            w.writerow([order, day.isoformat(), who, city, what, rng.randrange(1, 6), f"{rng.randrange(5, 400)}.{rng.randrange(100):02d}"])


def main() -> None:
    OUT.mkdir(exist_ok=True)
    pics = "/Library/User Pictures"
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        deck("board-update.pptx", [
            {"title": "Board update, November 2026", "bullets": ["Harbor Analytics", "Prepared for the quarterly board meeting"]},
            {"title": "Customer churn", "bullets": ["Monthly churn fell from 3.4% to 2.1%",
                                                   "Biggest loss: two retail clients in Rotterdam"]},
            {"title": "Warehouse automation", "bullets": ["Conveyor upgrade finished in Gdansk",
                                                         "Picking errors down 40% since August"]},
            {"picture": f"{pics}/Sports/Target.heic"},
            {"title": "Next steps", "bullets": ["Open a support desk in Montreal",
                                               "Renew the freight contract with Baltic Lines"]},
        ], tmp)
        deck("volcanoes.pptx", [
            {"title": "How volcanoes form", "bullets": ["Magma rises where tectonic plates pull apart",
                                                       "or where one plate dives under another"]},
            {"title": "Famous eruptions", "bullets": ["Vesuvius buried Pompeii in 79 AD",
                                                     "Krakatoa in 1883 was heard 4,800 km away"]},
            {"picture": f"{pics}/Flowers/Whiterose.heic"},
            {"title": "Our model volcano", "bullets": ["Baking soda, vinegar and red food colouring",
                                                      "in a cone of modelling clay"]},
        ], tmp)
        deck("messeauftritt.pptx", [
            {"title": "Messeauftritt Hannover 2027", "bullets": ["Planung für die Industriemesse im April"]},
            {"title": "Standplanung", "bullets": ["Stand mit 40 Quadratmetern in Halle 7",
                                                 "Zwei Bildschirme und eine Kaffeebar"]},
            {"picture": f"{pics}/Fun/Yin-Yang.heic"},
            {"title": "Kosten", "bullets": ["Standmiete 18.500 Euro", "Reise und Hotel 6.200 Euro"]},
        ], tmp)

    d = datetime.date
    guests = [["Name", "Side", "Table", "Meal", "RSVP"]]
    first = ["Oliver", "Mia", "Lucas", "Sofia", "Noah", "Emma", "Leo", "Ava", "Elias", "Ida", "Hugo", "Nora"]
    last = ["Berg", "Silva", "Novak", "Larsen", "Costa", "Weber", "Moreno", "Janssen", "Dubois", "Kowalczyk"]
    rng = random.Random(3)
    for i in range(36):
        guests.append([f"{rng.choice(first)} {rng.choice(last)}", rng.choice(["Bride", "Groom"]), rng.randrange(1, 9),
                       rng.choice(["Fish", "Beef", "Vegetarian"]), rng.choice(["Yes", "Yes", "No"])])
    guests.insert(20, ["Priya Raman", "Bride", 4, "Vegetarian", "Yes"])
    workbook("wedding.xlsx", {
        "Guests": guests,
        "Vendors": [["Vendor", "Service", "Deposit", "Due"],
                    ["Bloom & Twig", "Flowers", 450, d(2027, 3, 1)],
                    ["Lakeside Hall", "Venue", 3000, d(2027, 1, 15)],
                    ["DJ Marco", "Music", 300, d(2027, 4, 10)],
                    ["Sugar Loaf Bakery", "Cake", 520, d(2027, 4, 20)]],
    })
    schedule = [["Month", "Payment", "Interest", "Principal", "Balance"]]
    balance, rate = 24000.0, 0.049 / 12
    pay = round(balance * rate / (1 - (1 + rate) ** -48), 2)
    for m in range(1, 49):
        interest = round(balance * rate, 2)
        balance = round(balance - (pay - interest), 2)
        schedule.append([m, pay, interest, round(pay - interest, 2), max(balance, 0)])
    workbook("car-loan.xlsx", {
        "Summary": [["Loan amount", 24000], ["Interest rate", "4.9%"], ["Term", "48 months"],
                    ["Lender", "Riverbank Credit Union"], ["First payment", d(2026, 2, 1)]],
        "Schedule": schedule,
    })
    workbook("inventory.xlsx", {"Stock": inventory()})
    orders()
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    rain = [["Month", "Rainfall mm"]] + [[m, v] for m, v in zip(months, [78, 61, 55, 42, 50, 33, 21, 28, 47, 88, 102, 95])]
    workbook("rainfall.xlsx", {"Rain": rain}, chart="Rain")

    numbers("chess-club.numbers", {
        "Members": ("Club members", [["Name", "Rating", "Joined"], ["Tomasz Wierzbicki", 1840, 2019],
                                     ["Lena Hofmann", 1655, 2022], ["Arjun Mehta", 1720, 2021], ["Chloe Martin", 1510, 2024]]),
        "Tournament": ("Autumn tournament", [["Round", "White", "Black", "Result"],
                                             [1, "Lena Hofmann", "Tomasz Wierzbicki", "0-1"],
                                             [2, "Arjun Mehta", "Chloe Martin", "1-0"],
                                             [3, "Tomasz Wierzbicki", "Arjun Mehta", "draw"]]),
    })
    numbers("bike.numbers", {
        "Service": ("Bike service", [["Date", "Work", "Cost"], ["March", "tubeless tyres fitted", 120],
                                     ["May", "chain and cassette replaced", 85], ["August", "derailleur tuned", 25]]),
        "Rides": ("Long rides", [["Route", "Km"], ["Lake loop", 64], ["Hill climb to the radio mast", 41]]),
    })
    for f in sorted(OUT.iterdir()):
        print(f"{f.stat().st_size:8d}  {f.name}")


if __name__ == "__main__":
    main()
